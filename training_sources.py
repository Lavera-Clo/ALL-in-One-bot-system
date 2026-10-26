"""
More ways for a business to train its own bot, each feeding the training that already exists:

    website link   -> the pages' text      -> a document, trained like pasted text (training_pipeline)
    photos         -> one PDF of the photos -> a document, read by the AI vision readers (two models
                                               cross-check every number on scanned pages)
    FAQ sheet      -> question/answer rows  -> taught answers (learning_service.teach_many)

Nothing here decides what the bot says: the text goes through the same checked pipeline as a PDF.
"""

import asyncio
import csv
import io
import ipaddress
import logging
import re
import socket
from typing import List, Optional, Tuple
from urllib.parse import urldefrag, urljoin, urlparse

import httpx

logger = logging.getLogger("uvicorn.error")

# --- Website ------------------------------------------------------------------------------------

MAX_PAGES = 8
MAX_PAGE_BYTES = 2 * 1024 * 1024
MAX_SITE_CHARS = 120_000
MAX_REDIRECTS = 3
TIMEOUT = 15.0
# Pages most likely to hold what customers ask about are read first.
USEFUL = re.compile(r"price|pricing|rate|menu|product|shop|store|service|package|contact|deliver|shipping|"
                    r"return|refund|policy|faq|hours|location|book|room|order", re.I)
SKIP = re.compile(r"\.(jpg|jpeg|png|gif|webp|svg|pdf|zip|mp4|mp3|css|js|ico|xml)(\?|$)|/(login|signin|cart|checkout|"
                  r"account|wp-admin|wp-login|feed)\b|^mailto:|^tel:|^javascript:", re.I)


class SourceError(ValueError):
    """A problem with what the owner gave us, worded for them."""


def _public_host(host: str) -> None:
    """Refuse addresses inside a private network (this server's own, the cloud's metadata service...)."""
    if not host:
        raise SourceError("That isn't a website address.")
    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror:
        raise SourceError(f"Couldn't find the website {host}. Check the address.")
    for info in infos:
        ip = ipaddress.ip_address(info[4][0].split("%")[0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast or ip.is_unspecified:
            raise SourceError("That address points inside a private network, so it can't be read.")


def normalise_url(url: str) -> str:
    url = (url or "").strip()
    if not url:
        raise SourceError("Paste the website address.")
    if re.match(r"^[a-z][a-z0-9+.-]*://", url, re.I) and not re.match(r"^https?://", url, re.I):
        raise SourceError("Only website addresses (http or https) can be read.")
    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url
    parts = urlparse(url)
    host = parts.hostname or ""
    valid_host = re.fullmatch(r"[a-z0-9.-]+|[0-9a-f:]+", host, re.I) and ("." in host or ":" in host or host == "localhost")
    if parts.scheme not in ("http", "https") or not valid_host:
        raise SourceError("That isn't a website address.")
    return urldefrag(url)[0]


async def _get(client: httpx.AsyncClient, url: str) -> Tuple[str, str]:
    """(final url, html). Follows redirects one at a time so each hop is checked too."""
    for _ in range(MAX_REDIRECTS + 1):
        await asyncio.to_thread(_public_host, urlparse(url).hostname or "")
        async with client.stream("GET", url) as res:
            if res.status_code in (301, 302, 303, 307, 308) and res.headers.get("location"):
                url = urljoin(url, res.headers["location"])
                continue
            if res.status_code >= 400:
                raise SourceError(f"The website answered {res.status_code} for {url}.")
            if "html" not in (res.headers.get("content-type") or "").lower():
                raise SourceError(f"{url} isn't a web page.")
            body = b""
            async for chunk in res.aiter_bytes():
                body += chunk
                if len(body) > MAX_PAGE_BYTES:
                    break
            return url, body.decode(res.encoding or "utf-8", errors="replace")
    raise SourceError("The website redirects too many times.")


def page_text(html: str) -> Tuple[str, str]:
    """(title, readable text) of one page: no scripts, menus' repeated links collapse, tables as rows."""
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "iframe", "form", "template"]):
        tag.decompose()
    title = (soup.title.get_text(" ", strip=True) if soup.title else "")[:200]
    for row in soup.find_all("tr"):
        cells = [c.get_text(" ", strip=True) for c in row.find_all(["td", "th"])]
        row.replace_with(soup.new_string(" | ".join(c for c in cells if c) + "\n"))
    text = soup.get_text("\n", strip=True)
    lines, seen = [], set()
    for line in (re.sub(r"\s+", " ", l).strip() for l in text.split("\n")):
        if line and (len(line) > 40 or line not in seen):   # repeated short menu items once
            lines.append(line)
            seen.add(line)
    return title, "\n".join(lines)


def _links(base: str, html: str) -> List[str]:
    from bs4 import BeautifulSoup
    host = urlparse(base).hostname
    out = []
    for a in BeautifulSoup(html, "html.parser").find_all("a", href=True):
        href = a["href"].strip()
        if SKIP.search(href):
            continue
        url = urldefrag(urljoin(base, href))[0]
        if urlparse(url).hostname == host and urlparse(url).scheme in ("http", "https"):
            out.append(url)
    return out


async def read_website(url: str, max_pages: int = MAX_PAGES) -> dict:
    """{"url", "title", "text", "pages": [urls]} - the home page plus the most useful pages it links to."""
    url = normalise_url(url)
    pages: List[Tuple[str, str, str]] = []
    queue, seen = [url], set()
    headers = {"User-Agent": "Mozilla/5.0 (compatible; BusinessBotTrainer/1.0)", "Accept": "text/html"}
    async with httpx.AsyncClient(timeout=TIMEOUT, headers=headers, follow_redirects=False) as client:
        while queue and len(pages) < max_pages:
            next_url = queue.pop(0)
            if next_url in seen:
                continue
            seen.add(next_url)
            try:
                final, html = await _get(client, next_url)
            except SourceError:
                if not pages:
                    raise
                continue
            except httpx.HTTPError as exc:
                if not pages:
                    raise SourceError(f"Couldn't open {next_url} ({exc.__class__.__name__}).")
                continue
            seen.add(final)
            title, text = page_text(html)
            if text:
                pages.append((final, title, text))
            found = [u for u in _links(final, html) if u not in seen and u not in queue]
            queue += sorted(found, key=lambda u: 0 if USEFUL.search(u) else 1)
    if not pages:
        raise SourceError("No readable text on that website.")
    parts, total = [], 0
    for page_url, title, text in pages:
        block = f"## {title or page_url}\n(Source: {page_url})\n{text}"
        if total + len(block) > MAX_SITE_CHARS:
            block = block[: max(0, MAX_SITE_CHARS - total)]
        parts.append(block)
        total += len(block)
        if total >= MAX_SITE_CHARS:
            break
    return {"url": url, "title": pages[0][1] or urlparse(url).hostname, "text": "\n\n".join(parts),
            "pages": [p[0] for p in pages]}


# --- Photos ---------------------------------------------------------------------------------------

MAX_PHOTOS = 20
MAX_PHOTO_BYTES = 10 * 1024 * 1024
MAX_EDGE = 2000


def photos_to_pdf(photos: List[Tuple[bytes, str]]) -> bytes:
    """Photos of a menu, price board or product list -> one PDF, a page per photo."""
    from PIL import Image, ImageOps
    if not photos:
        raise SourceError("Choose at least one photo.")
    if len(photos) > MAX_PHOTOS:
        raise SourceError(f"Up to {MAX_PHOTOS} photos at a time.")
    pages = []
    for data, name in photos:
        if len(data) > MAX_PHOTO_BYTES:
            raise SourceError(f"{name} is over {MAX_PHOTO_BYTES // (1024 * 1024)} MB.")
        try:
            img = Image.open(io.BytesIO(data))
            img = ImageOps.exif_transpose(img).convert("RGB")     # phone photos: the right way up
        except Exception:  # noqa: BLE001
            raise SourceError(f"{name} isn't a photo this can read (use JPG or PNG).")
        img.thumbnail((MAX_EDGE, MAX_EDGE))
        pages.append(img)
    out = io.BytesIO()
    pages[0].save(out, format="PDF", save_all=True, append_images=pages[1:], resolution=150)
    return out.getvalue()


# --- FAQ spreadsheet ------------------------------------------------------------------------------

MAX_FAQ_ROWS = 500
_HEADER_Q = re.compile(r"^(q|question|questions|ask|query|ප්‍රශ්නය)$", re.I)


def _rows_from_csv(content: bytes) -> List[List[str]]:
    text = content.decode("utf-8-sig", errors="replace")
    lines = [l for l in text.splitlines() if l.strip()]
    # Tab or semicolon when every line has one (Excel saves ";" in many countries); else comma.
    delimiter = next((d for d in ("\t", ";") if lines and all(d in l for l in lines)), ",")
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _rows_from_xlsx(content: bytes) -> List[List[str]]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        raise SourceError("That spreadsheet couldn't be opened. Save it as .xlsx or .csv.")
    ws = wb.worksheets[0]
    return [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]


def parse_faq(content: bytes, filename: str) -> List[Tuple[str, str]]:
    """(question, answer) pairs: the first two filled columns of each row; a header row is skipped."""
    name = (filename or "").lower()
    typed = False
    if name.endswith((".xlsx", ".xlsm")):
        rows = _rows_from_xlsx(content)
    elif name.endswith((".csv", ".txt", ".tsv")):
        rows = _rows_from_csv(content)
        typed = True       # a hand-typed CSV often has unquoted commas in the answer
    else:
        raise SourceError("Upload a .csv or .xlsx file with questions in the first column and answers in the second.")
    pairs = []
    for i, row in enumerate(rows):
        cells = [str(c).strip() for c in row if str(c).strip()]
        if len(cells) < 2:
            continue
        q, a = cells[0], (", ".join(cells[1:]) if typed else cells[1])
        if i == 0 and _HEADER_Q.match(q):
            continue
        pairs.append((q, a))
        if len(pairs) >= MAX_FAQ_ROWS:
            break
    if not pairs:
        raise SourceError("No questions and answers found. Put each question in column A and its answer in column B.")
    return pairs
